"""1-4. 관리문서 작성

[+ ITEM 추가]를 누를 때마다 작업 블록(ITEM 선택 + 라인명 선택 + 공정마스터/문서 버튼)이 아래로 계속 추가된다.
각 블록에서 공정마스터 양식을 첨부하고, 마스터 정보로 공정특별특성 / 관리계획서 / 제조공정도 / PFMEA /
작업표준 엑셀을 자동 생성해 엑셀에서 작성한다. 엑셀에서 저장할 때마다 파일 변경을 감지해 문서별로 REV를 자동 기록한다.
"""
import os
import shutil
import uuid
from datetime import datetime
from itertools import islice, zip_longest
from pathlib import Path

from openpyxl import load_workbook
from PySide6.QtCore import Qt, QTimer, Signal
from PySide6.QtWidgets import (
    QAbstractItemView, QApplication, QComboBox, QFileDialog, QFrame, QHBoxLayout, QHeaderView,
    QLabel, QMessageBox, QPushButton, QScrollArea, QSplitter, QTableWidget,
    QTableWidgetItem, QVBoxLayout, QWidget,
)

from app.apqp_bridge import EngineMissingError, generate_document_from_master

POLL_INTERVAL_MS = 2000
COMPARE_MAX_ROWS = 5000
_CURRENT = object()

KINDS = [
    ("master", "공정마스터 양식"),
    ("special", "공정특별특성"),
    ("cp", "관리계획서"),
    ("flow", "제조공정도"),
    ("pfmea", "PFMEA"),
    ("wi", "작업표준"),
]
KIND_LABEL = dict(KINDS)
DERIVED_KINDS = [k for k, _ in KINDS[1:]]

CARD_STYLE = """
QFrame#docCard {
    background: #ffffff;
    border: 1px solid #e2e8f0;
    border-radius: 12px;
}
"""
COMBO_STYLE = "QComboBox { background: #f8fafc; border: 1px solid #cbd5e1; border-radius: 5px; padding: 5px 8px; color: #000000; font-size: 12px; min-width: 150px; }"
PRIMARY_BTN = """
QPushButton { background: #2563eb; color: #ffffff; font-size: 11px; font-weight: 600; padding: 8px 14px; border-radius: 6px; border: none; }
QPushButton:hover { background: #1d4ed8; }
QPushButton:disabled { background: #cbd5e1; }
"""
WRITE_BTN = PRIMARY_BTN.replace("#2563eb", "#059669").replace("#1d4ed8", "#047857")
DOC_BTN = """
QPushButton { background: #1e3a8a; color: #ffffff; font-size: 11px; font-weight: 600; padding: 8px 14px; border-radius: 6px; border: none; }
QPushButton:hover { background: #1e40af; }
QPushButton:disabled { background: #cbd5e1; }
"""
SUB_BTN = """
QPushButton { background: #f1f5f9; color: #000000; font-size: 11px; font-weight: 600; padding: 6px 12px; border-radius: 6px; border: 1px solid #cbd5e1; }
QPushButton:hover { background: #e2e8f0; }
QPushButton:disabled { color: #94a3b8; }
"""
ADD_BLOCK_BTN = """
QPushButton { background: #2563eb; color: #ffffff; font-size: 12px; font-weight: 700; padding: 8px 18px; border-radius: 8px; border: none; }
QPushButton:hover { background: #1d4ed8; }
"""
REMOVE_BTN = """
QPushButton { background: transparent; color: #64748b; font-size: 11px; border: none; padding: 2px 6px; }
QPushButton:hover { color: #dc2626; }
"""
FIELD_LABEL = "color: #64748b; font-size: 10px; font-weight: 600;"
TABLE_STYLE = "QTableWidget { gridline-color: #e2e8f0; color: #000000; font-size: 12px; } QHeaderView::section { background: #f1f5f9; color: #475569; border: none; border-bottom: 1px solid #e2e8f0; padding: 4px; }"


def count_changes(old_path, new_path):
    """두 엑셀 파일을 비교해 값이 달라진 셀 개수를 센다 (실패하면 0)."""
    old_wb = new_wb = None
    total = 0
    try:
        old_wb = load_workbook(old_path, read_only=True)
        new_wb = load_workbook(new_path, read_only=True)
        for name in new_wb.sheetnames:
            new_rows = islice(new_wb[name].iter_rows(values_only=True), COMPARE_MAX_ROWS)
            old_rows = islice(old_wb[name].iter_rows(values_only=True), COMPARE_MAX_ROWS) if name in old_wb.sheetnames else iter(())
            for old_row, new_row in zip_longest(old_rows, new_rows, fillvalue=()):
                for old_val, new_val in zip_longest(old_row or (), new_row or (), fillvalue=None):
                    if old_val != new_val:
                        total += 1
    except Exception:
        return 0
    finally:
        for wb in (old_wb, new_wb):
            if wb:
                wb.close()
    return total


class EntryBlock(QFrame):
    """ITEM 선택 + 라인명 선택 + 공정마스터/문서 버튼 한 세트."""

    selectionChanged = Signal(object)
    activateRequested = Signal(object)
    removeRequested = Signal(object)
    attachRequested = Signal(object)
    masterWriteRequested = Signal(object)
    docRequested = Signal(object, str)

    def __init__(self, number, item_names, line_names, item="", line="", parent=None):
        super().__init__(parent)
        self.setObjectName("entryBlock")
        self.set_active(False)

        col = QVBoxLayout(self)
        col.setContentsMargins(14, 8, 14, 10)
        col.setSpacing(6)

        header = QHBoxLayout()
        self.title = QLabel("")
        self.title.setStyleSheet("color: #000000; font-weight: 700; font-size: 12px;")
        header.addWidget(self.title)
        header.addStretch(1)
        self.remove_btn = QPushButton("삭제")
        self.remove_btn.setStyleSheet(REMOVE_BTN)
        self.remove_btn.setToolTip("이 작업 줄을 목록에서 지웁니다 (저장된 문서 파일은 유지됩니다)")
        header.addWidget(self.remove_btn)
        col.addLayout(header)

        row1 = QHBoxLayout()
        row1.setSpacing(10)
        self.item_combo = QComboBox()
        self.item_combo.setEditable(True)
        self.item_combo.setStyleSheet(COMBO_STYLE)
        self.item_combo.lineEdit().setPlaceholderText("ITEM 선택/입력")
        self.line_combo = QComboBox()
        self.line_combo.setEditable(True)
        self.line_combo.setStyleSheet(COMBO_STYLE)
        self.line_combo.lineEdit().setPlaceholderText("라인명 선택/입력")
        for label_text, widget in (("ITEM 선택", self.item_combo), ("라인명 선택", self.line_combo)):
            box = QVBoxLayout()
            box.setSpacing(2)
            lab = QLabel(label_text)
            lab.setStyleSheet(FIELD_LABEL)
            box.addWidget(lab)
            box.addWidget(widget)
            row1.addLayout(box)
        row1.addStretch(1)
        col.addLayout(row1)

        row2 = QHBoxLayout()
        row2.setSpacing(8)
        self.attach_btn = QPushButton("공정마스터 양식 첨부")
        self.attach_btn.setStyleSheet(PRIMARY_BTN)
        row2.addWidget(self.attach_btn)
        self.master_write_btn = QPushButton("마스터 작성")
        self.master_write_btn.setStyleSheet(WRITE_BTN)
        self.master_write_btn.setEnabled(False)
        row2.addWidget(self.master_write_btn)
        separator = QFrame()
        separator.setFrameShape(QFrame.VLine)
        separator.setStyleSheet("color: #cbd5e1;")
        row2.addWidget(separator)
        self.doc_buttons = {}
        for kind in DERIVED_KINDS:
            btn = QPushButton(KIND_LABEL[kind])
            btn.setStyleSheet(DOC_BTN)
            btn.setEnabled(False)
            btn.clicked.connect(lambda _c=False, k=kind: self.docRequested.emit(self, k))
            row2.addWidget(btn)
            self.doc_buttons[kind] = btn
        row2.addStretch(1)
        col.addLayout(row2)

        self.set_item_names(item_names, item)
        self.set_line_names(line_names, line)
        self.set_number(number)

        self.remove_btn.clicked.connect(lambda: self.removeRequested.emit(self))
        self.attach_btn.clicked.connect(lambda: self.attachRequested.emit(self))
        self.master_write_btn.clicked.connect(lambda: self.masterWriteRequested.emit(self))
        for combo in (self.item_combo, self.line_combo):
            combo.activated.connect(lambda _i: self.selectionChanged.emit(self))
            combo.lineEdit().editingFinished.connect(lambda: self.selectionChanged.emit(self))

    def set_number(self, number):
        self.title.setText(f"ITEM 작업 {number}")

    def set_item_names(self, names, keep=None):
        keep = self.item_combo.currentText() if keep is None else keep
        self.item_combo.blockSignals(True)
        self.item_combo.clear()
        self.item_combo.addItems(names)
        self.item_combo.setEditText(keep)
        self.item_combo.blockSignals(False)

    def set_line_names(self, names, keep=None):
        keep = self.line_combo.currentText() if keep is None else keep
        self.line_combo.blockSignals(True)
        self.line_combo.clear()
        self.line_combo.addItems(names)
        self.line_combo.setEditText(keep)
        self.line_combo.blockSignals(False)

    def item_text(self):
        return self.item_combo.currentText().strip()

    def line_text(self):
        return self.line_combo.currentText().strip()

    def base(self):
        item, line = self.item_text(), self.line_text()
        return f"{item}||{line}" if item and line else None

    def set_master_ready(self, ready):
        self.master_write_btn.setEnabled(ready)
        for btn in self.doc_buttons.values():
            btn.setEnabled(ready)

    def set_active(self, active):
        self._active = active
        border = "2px solid #2563eb" if active else "1px solid #e2e8f0"
        self.setStyleSheet(f"QFrame#entryBlock {{ background: #ffffff; border: {border}; border-radius: 12px; }}")

    def mousePressEvent(self, event):
        self.activateRequested.emit(self)
        super().mousePressEvent(event)


class DocScreen(QWidget):
    itemsChanged = Signal()

    def __init__(self, state, parent=None):
        super().__init__(parent)
        self.state = state
        self.blocks = []
        self.active_block = None
        self.active_kind = "master"
        self._current_base = None
        self._pending = {}

        self.poll_timer = QTimer(self)
        self.poll_timer.setInterval(POLL_INTERVAL_MS)
        self.poll_timer.timeout.connect(self._poll)
        self.poll_timer.start()

        outer = QVBoxLayout(self)
        outer.setContentsMargins(4, 8, 4, 8)
        outer.setSpacing(10)

        top = QHBoxLayout()
        self.add_block_btn = QPushButton("+ ITEM 추가")
        self.add_block_btn.setStyleSheet(ADD_BLOCK_BTN)
        self.add_block_btn.setToolTip("누를 때마다 아래에 ITEM 작업 줄이 추가됩니다")
        self.add_block_btn.clicked.connect(lambda: self._add_block())
        top.addWidget(self.add_block_btn)
        self.status_label = QLabel("")
        self.status_label.setStyleSheet("color: #059669; font-size: 11px; font-weight: 600;")
        top.addWidget(self.status_label, 1)
        outer.addLayout(top)

        vsplit = QSplitter(Qt.Vertical)
        vsplit.setChildrenCollapsible(False)

        self.blocks_scroll = QScrollArea()
        self.blocks_scroll.setWidgetResizable(True)
        self.blocks_scroll.setStyleSheet("QScrollArea { border: none; background: transparent; }")
        self.blocks_container = QWidget()
        self.blocks_layout = QVBoxLayout(self.blocks_container)
        self.blocks_layout.setContentsMargins(0, 0, 4, 0)
        self.blocks_layout.setSpacing(10)
        self.blocks_layout.addStretch(1)
        self.blocks_scroll.setWidget(self.blocks_container)
        vsplit.addWidget(self.blocks_scroll)

        hsplit = QSplitter(Qt.Horizontal)
        hsplit.setChildrenCollapsible(False)
        hsplit.addWidget(self._build_status_card())
        hsplit.addWidget(self._build_rev_card())
        hsplit.setStretchFactor(0, 3)
        hsplit.setStretchFactor(1, 2)
        vsplit.addWidget(hsplit)
        vsplit.setStretchFactor(0, 3)
        vsplit.setStretchFactor(1, 2)
        outer.addWidget(vsplit, 1)

    # ---------------- 하단 패널 ----------------
    def _build_status_card(self):
        card = QFrame()
        card.setObjectName("docCard")
        card.setStyleSheet(CARD_STYLE)
        layout = QVBoxLayout(card)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(8)

        self.status_title = QLabel("문서 현황")
        self.status_title.setStyleSheet("font-weight: 700; color: #000000; font-size: 13px;")
        layout.addWidget(self.status_title)

        self.status_table = QTableWidget(len(KINDS), 4)
        self.status_table.setHorizontalHeaderLabels(["문서", "상태", "최신 REV", "최종 저장"])
        self.status_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.status_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.status_table.setSelectionMode(QAbstractItemView.SingleSelection)
        self.status_table.verticalHeader().hide()
        self.status_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        self.status_table.horizontalHeader().setStretchLastSection(True)
        self.status_table.setStyleSheet(TABLE_STYLE)
        self.status_table.cellClicked.connect(self._on_status_row_clicked)
        layout.addWidget(self.status_table, 1)

        hint = QLabel("파란 테두리로 표시된 ITEM 작업 줄의 문서 현황입니다. 행을 클릭하면 오른쪽에 해당 문서의 REV 이력이 표시됩니다.")
        hint.setWordWrap(True)
        hint.setStyleSheet("color: #64748b; font-size: 10px;")
        layout.addWidget(hint)
        return card

    def _build_rev_card(self):
        card = QFrame()
        card.setObjectName("docCard")
        card.setStyleSheet(CARD_STYLE)
        card.setMinimumWidth(320)
        layout = QVBoxLayout(card)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(8)

        self.rev_title = QLabel("REV 이력")
        self.rev_title.setStyleSheet("font-weight: 700; color: #000000; font-size: 13px;")
        layout.addWidget(self.rev_title)

        self.rev_table = QTableWidget(0, 4)
        self.rev_table.setHorizontalHeaderLabels(["REV", "저장 일시", "변경", "메모"])
        self.rev_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.rev_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.rev_table.setSelectionMode(QAbstractItemView.SingleSelection)
        self.rev_table.verticalHeader().hide()
        self.rev_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        self.rev_table.horizontalHeader().setStretchLastSection(True)
        self.rev_table.setStyleSheet(TABLE_STYLE.replace("12px", "11px"))
        layout.addWidget(self.rev_table, 1)

        self.restore_btn = QPushButton("선택한 REV로 되돌리기")
        self.restore_btn.setStyleSheet(SUB_BTN)
        self.restore_btn.clicked.connect(self._restore_selected_rev)
        layout.addWidget(self.restore_btn)

        hint = QLabel("엑셀에서 저장할 때마다 자동으로 새 REV가 기록됩니다.")
        hint.setWordWrap(True)
        hint.setStyleSheet("color: #64748b; font-size: 10px;")
        layout.addWidget(hint)
        return card

    # ---------------- 작업 블록(ITEM 줄) 관리 ----------------
    def refresh(self):
        """화면에 들어올 때마다 저장된 작업 줄을 다시 구성한다."""
        self._poll(force=True)
        old_index = self.blocks.index(self.active_block) if self.active_block in self.blocks else 0

        for block in self.blocks:
            self.blocks_layout.removeWidget(block)
            block.deleteLater()
        self.blocks = []
        self.active_block = None

        entries = self.state.doc_entries or [{"item": self.state.current_item, "line": ""}]
        for entry in entries:
            self._add_block(entry.get("item", ""), entry.get("line", ""), save=False, scroll=False)
        self._set_active(self.blocks[min(old_index, len(self.blocks) - 1)], force=True)

    def _add_block(self, item=None, line="", save=True, scroll=True):
        number = len(self.blocks) + 1
        if item is None:
            item = self.active_block.item_text() if self.active_block else self.state.current_item
        block = EntryBlock(number, self.state.item_choices(), self.state.line_names, item, line, self.blocks_container)
        block.selectionChanged.connect(self._on_block_changed)
        block.activateRequested.connect(self._set_active)
        block.removeRequested.connect(self._remove_block)
        block.attachRequested.connect(self._attach_master)
        block.masterWriteRequested.connect(lambda b: self._open_in_excel(b, "master"))
        block.docRequested.connect(self._open_generated)

        self.blocks_layout.insertWidget(self.blocks_layout.count() - 1, block)
        self.blocks.append(block)
        if save:
            self._save_entries()
            self._set_active(block)
            if scroll:
                QTimer.singleShot(0, lambda: self.blocks_scroll.verticalScrollBar().setValue(
                    self.blocks_scroll.verticalScrollBar().maximum()))
        return block

    def _remove_block(self, block):
        if len(self.blocks) <= 1:
            QMessageBox.information(self, "알림", "마지막 작업 줄은 삭제할 수 없습니다.")
            return
        self._poll(force=True)
        index = self.blocks.index(block)
        self.blocks.remove(block)
        self.blocks_layout.removeWidget(block)
        block.deleteLater()
        for number, remaining in enumerate(self.blocks, start=1):
            remaining.set_number(number)
        self._save_entries()
        if block is self.active_block:
            self.active_block = None
            self._set_active(self.blocks[min(index, len(self.blocks) - 1)], force=True)
        else:
            self._refresh_all()

    def _save_entries(self):
        self.state.data["doc_entries"] = [{"item": b.item_text(), "line": b.line_text()} for b in self.blocks]
        self.state.save()

    def _on_block_changed(self, block):
        item, line = block.item_text(), block.line_text()
        if item and item not in self.state.item_order:
            previous_current = self.state.current_item
            if self.state.add_item(item):
                self.state.current_item = previous_current
                self.state.save()
                self.itemsChanged.emit()
                for b in self.blocks:
                    b.set_item_names(self.state.item_choices())
        if line and line not in self.state.line_names:
            self.state.add_line_name(line)
            for b in self.blocks:
                b.set_line_names(self.state.line_names)
        self._save_entries()
        self._set_active(block)

    def _set_active(self, block, force=False):
        if block is None:
            return
        if not force and block is self.active_block and block.base() == self._current_base:
            return
        self._poll(force=True)
        self.active_block = block
        for b in self.blocks:
            b.set_active(b is block)
        self._current_base = block.base()
        self._refresh_all()

    # ---------------- 문서 조회 ----------------
    def _doc_key(self, kind, base=_CURRENT):
        base = self._current_base if base is _CURRENT else base
        return f"{base}||{kind}" if base else None

    def _doc(self, kind, base=_CURRENT):
        key = self._doc_key(kind, base)
        return self.state.docs.get(key) if key else None

    def _working_path(self, doc):
        return self.state.doc_dir(doc) / f"working{doc.get('ext', '.xlsx')}"

    def _doc_ready(self, kind, base=_CURRENT):
        doc = self._doc(kind, base)
        return bool(doc and self._working_path(doc).exists())

    def _latest_rev(self, doc):
        revs = doc.get("revs", []) if doc else []
        return revs[-1] if revs else None

    def _migrate_legacy_doc(self, base):
        """이전 버전(ITEM||라인명 키로 저장된 단일 문서)을 공정마스터 문서로 옮긴다."""
        legacy = self.state.docs.get(base) if base else None
        new_key = self._doc_key("master", base)
        if legacy and new_key and new_key not in self.state.docs:
            legacy["key"] = new_key
            legacy["kind"] = "master"
            self.state.docs[new_key] = self.state.docs.pop(base)
            self.state.save()

    # ---------------- 화면 갱신 ----------------
    def _refresh_all(self):
        self._pending.clear()
        for block in self.blocks:
            base = block.base()
            self._migrate_legacy_doc(base)
            for kind, _label in KINDS:
                doc = self._doc(kind, base)
                if doc and doc.get("mtime") is None and self._doc_ready(kind, base):
                    doc["mtime"] = self._working_path(doc).stat().st_mtime
            block.set_master_ready(self._doc_ready("master", base))

        self._refresh_status_table()
        self._refresh_rev_table()
        if not self._current_base:
            self._set_status("ITEM과 라인명을 선택(입력)해주세요.", warn=True)
        elif not self._doc_ready("master"):
            self._set_status("공정마스터 양식을 먼저 첨부해주세요.", warn=True)
        else:
            self._set_status("")

    def _refresh_status_table(self):
        if self.active_block:
            number = self.blocks.index(self.active_block) + 1
            self.status_title.setText(f"문서 현황 - ITEM 작업 {number}")
        master_rev = self._latest_rev(self._doc("master"))
        for row, (kind, label) in enumerate(KINDS):
            doc = self._doc(kind)
            ready = self._doc_ready(kind)
            latest = self._latest_rev(doc) if ready else None

            if kind == "master":
                state_text = "첨부됨" if ready else "미첨부"
            elif not ready:
                state_text = "미생성"
            elif master_rev and master_rev["rev"] > doc.get("master_rev", -1):
                state_text = "마스터 변경됨 (재생성 필요)"
            else:
                state_text = "생성됨"

            values = [
                label, state_text,
                f"REV {latest['rev']}" if latest else "-",
                latest["time"].replace("T", " ") if latest else "-",
            ]
            for col, text in enumerate(values):
                item = QTableWidgetItem(text)
                item.setData(Qt.UserRole, kind)
                self.status_table.setItem(row, col, item)
        self.status_table.selectRow([k for k, _ in KINDS].index(self.active_kind))

    def _on_status_row_clicked(self, row, _col):
        self.active_kind = KINDS[row][0]
        self._refresh_rev_table()

    def _set_status(self, text, warn=False, error=False):
        color = "#dc2626" if error else "#d97706" if warn else "#059669"
        self.status_label.setStyleSheet(f"color: {color}; font-size: 11px; font-weight: 600;")
        self.status_label.setText(text)

    # ---------------- 공정마스터 첨부 / 작성 ----------------
    def _attach_master(self, block):
        self._set_active(block)
        if not self._current_base:
            QMessageBox.warning(self, "알림", "먼저 ITEM과 라인명을 선택(입력)해주세요.")
            return
        path, _ = QFileDialog.getOpenFileName(self, "공정마스터 양식 선택", "", "엑셀 파일 (*.xlsx *.xlsm)")
        if not path:
            return
        ext = Path(path).suffix.lower()
        try:
            load_workbook(path, keep_vba=ext == ".xlsm").close()
        except Exception as e:
            QMessageBox.critical(self, "오류", f"엑셀 파일을 열 수 없습니다.\n(.xlsx / .xlsm 형식만 지원합니다)\n{e}")
            return

        self._poll(force=True)
        doc = self._doc("master")
        replacing = self._doc_ready("master")
        if replacing:
            reply = QMessageBox.question(
                self, "마스터 교체",
                "이미 첨부된 공정마스터가 있습니다. 새 파일로 교체할까요?\n(기존 REV 이력은 유지되고, 새 REV로 기록됩니다)",
            )
            if reply != QMessageBox.Yes:
                return
        if not doc:
            doc = {"key": self._doc_key("master"), "kind": "master", "dir": uuid.uuid4().hex, "revs": []}
            self.state.docs[doc["key"]] = doc

        self.state.doc_dir(doc).mkdir(parents=True, exist_ok=True)
        old_working = self._working_path(doc)
        doc["ext"] = ext
        try:
            if old_working.exists() and old_working != self._working_path(doc):
                old_working.unlink()
            shutil.copyfile(path, self._working_path(doc))
        except PermissionError:
            QMessageBox.warning(self, "알림", "엑셀에서 이 문서가 열려 있습니다. 엑셀을 닫은 뒤 다시 시도해주세요.")
            return
        doc["source_name"] = Path(path).name
        self._record_revision(doc, note=("새 파일 첨부: " if replacing else "원본 첨부: ") + Path(path).name)
        doc["mtime"] = self._working_path(doc).stat().st_mtime
        self.state.save()
        self.active_kind = "master"
        self._refresh_all()

    def _open_in_excel(self, block, kind):
        self._set_active(block)
        doc = self._doc(kind)
        if not doc or not self._doc_ready(kind):
            QMessageBox.warning(self, "알림", "문서 파일을 찾을 수 없습니다.")
            return False
        try:
            os.startfile(str(self._working_path(doc)))
        except OSError as e:
            QMessageBox.critical(self, "오류", f"엑셀을 실행할 수 없습니다.\n{e}")
            return False
        self.active_kind = kind
        self._refresh_status_table()
        self._refresh_rev_table()
        self._set_status(f"엑셀에서 {KIND_LABEL[kind]} 작성 중... (저장하면 자동으로 REV가 기록됩니다)", warn=True)
        return True

    # ---------------- 마스터 -> 문서 생성 ----------------
    def _open_generated(self, block, kind):
        self._set_active(block)
        master = self._doc("master")
        if not master or not self._doc_ready("master"):
            QMessageBox.warning(self, "알림", "공정마스터 양식을 먼저 첨부해주세요.")
            return

        doc = self._doc(kind)
        master_rev = self._latest_rev(master)["rev"]
        need_generate = not self._doc_ready(kind)

        if not need_generate and master_rev > doc.get("master_rev", -1):
            box = QMessageBox(self)
            box.setWindowTitle("마스터 변경됨")
            box.setText(f"{KIND_LABEL[kind]}을(를) 만든 뒤 공정마스터가 변경되었습니다.\n다시 생성할까요?\n(다시 생성하면 현재 파일은 마스터 기준으로 새로 만들어지고, 이전 작성 내용은 REV 이력에 남습니다)")
            regenerate = box.addButton("다시 생성", QMessageBox.AcceptRole)
            keep = box.addButton("기존 파일 열기", QMessageBox.ActionRole)
            box.addButton("취소", QMessageBox.RejectRole)
            box.exec()
            clicked = box.clickedButton()
            if clicked is regenerate:
                need_generate = True
            elif clicked is not keep:
                return

        if need_generate and not self._generate(kind):
            return
        self._open_in_excel(block, kind)

    def _generate(self, kind):
        master = self._doc("master")
        master_rev = self._latest_rev(master)["rev"]
        doc = self._doc(kind)
        existed = bool(doc and self._doc_ready(kind))
        if not doc:
            doc = {"key": self._doc_key(kind), "kind": kind, "dir": uuid.uuid4().hex, "ext": ".xlsm", "revs": []}
            self.state.docs[doc["key"]] = doc
        self.state.doc_dir(doc).mkdir(parents=True, exist_ok=True)

        old_working = self._working_path(doc)
        doc["ext"] = ".xlsm"
        target = self._working_path(doc)
        self._set_status(f"{KIND_LABEL[kind]} 생성 중...", warn=True)
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            count, _skipped, std_source = generate_document_from_master(self._working_path(master), kind, target)
            if old_working.exists() and old_working != target:
                old_working.unlink()
        except EngineMissingError as e:
            QMessageBox.critical(self, "오류", str(e))
            return False
        except PermissionError:
            QMessageBox.warning(self, "알림", f"엑셀에서 {KIND_LABEL[kind]}이(가) 열려 있습니다. 엑셀을 닫은 뒤 다시 시도해주세요.")
            return False
        except Exception as e:
            QMessageBox.critical(self, "오류", f"문서를 생성하는 중 오류가 발생했습니다.\n{e}")
            return False
        finally:
            QApplication.restoreOverrideCursor()

        doc["source_name"] = f"{KIND_LABEL[kind]}.xlsm"
        doc["master_rev"] = master_rev
        note = f"마스터 REV {master_rev} 기준 " + ("재생성" if existed else "자동 생성") + f" (공정 {count}개, {std_source})"
        self._record_revision(doc, note=note)
        doc["mtime"] = self._working_path(doc).stat().st_mtime
        self.state.save()
        return True

    # ---------------- 저장 감지 -> REV 기록 ----------------
    def flush(self):
        """엑셀에서 저장된 변경이 아직 기록되지 않았다면 즉시 REV로 남긴다."""
        self._poll(force=True)

    def _poll(self, force=False):
        bases = []
        for block in self.blocks:
            base = block.base()
            if base and base not in bases:
                bases.append(base)

        captured = False
        editing = []
        for base in bases:
            for kind, label in KINDS:
                doc = self._doc(kind, base)
                if not doc or not self._doc_ready(kind, base):
                    continue
                path = self._working_path(doc)
                mtime = path.stat().st_mtime
                known = doc.get("mtime")
                if known is None:
                    doc["mtime"] = mtime
                elif mtime != known:
                    if force or self._pending.get(doc["dir"]) == mtime:
                        if self._capture_revision(doc, mtime):
                            captured = True
                            self._pending.pop(doc["dir"], None)
                    else:
                        self._pending[doc["dir"]] = mtime
                        self._set_status(f"{label} 엑셀 저장 감지... REV 기록 중", warn=True)
                if path.with_name("~$" + path.name).exists():
                    editing.append(label)

        if captured:
            self._refresh_status_table()
            self._refresh_rev_table()
        elif editing:
            self._set_status("엑셀에서 작성 중: " + ", ".join(editing) + " (저장하면 자동으로 REV가 기록됩니다)", warn=True)

    def _capture_revision(self, doc, mtime):
        previous_info = self._latest_rev(doc)
        previous = self.state.doc_dir(doc) / previous_info["file"] if previous_info else None
        try:
            changes = count_changes(previous, self._working_path(doc)) if previous and previous.exists() else 0
            rev = self._record_revision(doc, changes=changes)
        except PermissionError:
            return False
        except OSError as e:
            self._set_status(f"REV 기록 실패: {e}", error=True)
            return False
        doc["mtime"] = mtime
        self.state.save()
        self._set_status(f"{KIND_LABEL[doc['kind']]} 엑셀 저장 감지 · REV {rev} 자동 기록됨 · {datetime.now().strftime('%H:%M:%S')}")
        return True

    def _record_revision(self, doc, note="", changes=0):
        revs = doc.setdefault("revs", [])
        rev = (revs[-1]["rev"] + 1) if revs else 0
        rev_file = f"rev_{rev:03d}{doc.get('ext', '.xlsx')}"
        shutil.copyfile(self._working_path(doc), self.state.doc_dir(doc) / rev_file)
        revs.append({
            "rev": rev,
            "time": datetime.now().isoformat(timespec="seconds"),
            "file": rev_file,
            "changes": changes,
            "note": note,
        })
        self.state.save()
        return rev

    # ---------------- REV 이력 ----------------
    def _refresh_rev_table(self):
        doc = self._doc(self.active_kind)
        self.rev_title.setText(f"REV 이력 - {KIND_LABEL[self.active_kind]}")
        revs = list(reversed(doc.get("revs", []))) if doc else []
        self.rev_table.setRowCount(len(revs))
        for row, rev in enumerate(revs):
            values = [f"REV {rev['rev']}", rev["time"].replace("T", " "), str(rev.get("changes", 0)), rev.get("note", "")]
            for col, text in enumerate(values):
                item = QTableWidgetItem(text)
                item.setData(Qt.UserRole, rev["rev"])
                self.rev_table.setItem(row, col, item)
        self.restore_btn.setEnabled(bool(revs))

    def _restore_selected_rev(self):
        doc = self._doc(self.active_kind)
        if not doc:
            return
        selected = self.rev_table.selectedItems()
        if not selected:
            QMessageBox.information(self, "알림", "되돌릴 REV를 목록에서 먼저 선택해주세요.")
            return
        rev_no = selected[0].data(Qt.UserRole)
        rev = next((r for r in doc["revs"] if r["rev"] == rev_no), None)
        rev_path = self.state.doc_dir(doc) / rev["file"] if rev else None
        if not rev_path or not rev_path.exists():
            QMessageBox.warning(self, "알림", "해당 REV 파일을 찾을 수 없습니다.")
            return
        reply = QMessageBox.question(
            self, "REV 되돌리기",
            f"{KIND_LABEL[self.active_kind]}을(를) REV {rev_no} 내용으로 되돌릴까요?\n(현재 내용은 이력에 남고, 되돌린 결과가 새 REV로 기록됩니다)",
        )
        if reply != QMessageBox.Yes:
            return
        self._poll(force=True)
        old_working = self._working_path(doc)
        doc["ext"] = rev_path.suffix.lower()
        try:
            if old_working.exists() and old_working != self._working_path(doc):
                old_working.unlink()
            shutil.copyfile(rev_path, self._working_path(doc))
        except PermissionError:
            QMessageBox.warning(self, "알림", "엑셀에서 이 문서가 열려 있습니다. 엑셀을 닫은 뒤 다시 시도해주세요.")
            return
        self._record_revision(doc, note=f"REV {rev_no}에서 복원")
        doc["mtime"] = self._working_path(doc).stat().st_mtime
        self.state.save()
        self._refresh_all()
