"""공정마스터(.xlsm, '1.1MASTER' 시트)를 읽어 APQP 생성 엔진(engine/)으로 회사 양식 5종을 만든다.

엔진은 회사 양식(.xlsm)을 XML 수준에서 채우기 때문에 숨은 DataSheet(사내 서버 매핑용), 매크로, 도형이
원본 그대로 유지된다.
"""
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ENGINE_DIR = Path(__file__).resolve().parent.parent / "engine"
MASTER_SHEET = "1.1MASTER"
FIRST_PROCESS_ROW = 14
LAST_ROW = 400
BLOCK_ROWS = 9

KIND_TO_ENGINE_DOC = {
    "special": "공정특별특성",
    "cp": "관리계획서",
    "flow": "제조공정도",
    "pfmea": "PFMEA",
    "wi": "작업표준",
}


class EngineMissingError(RuntimeError):
    pass


def engine_available():
    return (ENGINE_DIR / "make.py").exists() and (ENGINE_DIR / "문서생성기.xlsm").exists() and (ENGINE_DIR / "templates").is_dir()


def _load_engine():
    if not engine_available():
        raise EngineMissingError(
            f"문서 생성 엔진을 찾을 수 없습니다.\n'{ENGINE_DIR}' 폴더에 core, templates, 문서생성기.xlsm, make.py가 있어야 합니다."
        )
    path = str(ENGINE_DIR)
    if path not in sys.path:
        sys.path.insert(0, path)
    import make  # noqa: E402
    from core import builders  # noqa: E402
    from core.masterdb import MasterDB  # noqa: E402
    return make, builders, MasterDB


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number_text(value):
    if isinstance(value, bool) or value is None:
        return ""
    if isinstance(value, (int, float)):
        return f"{value:g}"
    return _text(value)


def read_master_project(path):
    """1.1MASTER 시트에서 엔진 입력(project dict)을 만든다: {info, lines, rows}."""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        if MASTER_SHEET not in wb.sheetnames:
            raise ValueError(f"'{MASTER_SHEET}' 시트가 없습니다. 공정마스터 양식(R10.x)이 맞는지 확인해주세요.")
        grid = [list(r) for r in wb[MASTER_SHEET].iter_rows(min_row=1, max_row=LAST_ROW, max_col=12, values_only=True)]
    finally:
        wb.close()

    def cell(row, col):
        if row - 1 < len(grid) and col - 1 < len(grid[row - 1]):
            return grid[row - 1][col - 1]
        return None

    def clean(value):
        text = _text(value)
        return "" if text in ("-", "None") else text

    eo_values = [v for v in (clean(cell(8, 7)), clean(cell(10, 7))) if v]
    info = {
        "고객": clean(cell(4, 4)),
        "단계": clean(cell(4, 9)) or "P1",
        "품명": clean(cell(6, 7)),
        "Rev": clean(cell(6, 9)) or "0",
        "차종": clean(cell(8, 4)),
        "부서": clean(cell(8, 9)),
        "품번": clean(cell(10, 4)),
        "작성자": clean(cell(10, 9)),
        "EO": ", ".join(dict.fromkeys(eo_values)),
    }

    lines, rows = [], []
    for r in range(FIRST_PROCESS_ROW, min(len(grid), LAST_ROW) + 1):
        process_no = clean(cell(r, 4))
        if not process_no:
            continue
        line_name = clean(cell(r, 2))
        kind = clean(cell(r, 3))
        mark = "용접" if "용접" in kind else "조립"
        if line_name and not any(l["라인명"] == line_name for l in lines):
            lines.append({"구분": kind or mark, "표기": mark, "라인명": line_name, "적용": mark, "CT": "", "인원": ""})

        total = cell(r + 3, 12)
        for l in lines:
            if l["라인명"] == line_name and isinstance(total, (int, float)):
                l["CT"] = max(float(l["CT"] or 0), float(total))

        equipment = clean(cell(r + 7, 12))
        rows.append({
            "라인명": line_name or (lines[-1]["라인명"] if lines else ""),
            "공정번호": process_no,
            "공정Code": clean(cell(r, 5)),
            "공정명": clean(cell(r, 6)),
            "공정내용": clean(cell(r, 7)),
            "적용": mark,
            "설비명": equipment,
            "설비번호": "",
            "P1": "", "P2": "", "P3": "",
            "CT_L": _number_text(cell(r, 12)),
            "CT_W": _number_text(cell(r + 1, 12)),
            "CT_U": _number_text(cell(r + 2, 12)),
        })

    for l in lines:
        l["CT"] = _number_text(l["CT"]) if l["CT"] != "" else ""
    return {"info": info, "lines": lines, "rows": rows}


_db_cache = {}
EMBEDDED_SHEETS = {"STD_공정", "STD_관리항목", "STD_FMEA"}


def _has_embedded_std(master_path):
    wb = load_workbook(master_path, read_only=True)
    try:
        return EMBEDDED_SHEETS <= set(wb.sheetnames)
    finally:
        wb.close()


def _get_db(master_db_cls, master_path):
    """마스터에 고정값(STD 시트)이 내장돼 있으면 그것을, 없으면 엔진 기본 표준정보를 쓴다."""
    master_path = Path(master_path)
    embedded = _has_embedded_std(master_path)
    source = master_path if embedded else ENGINE_DIR / "문서생성기.xlsm"
    key = (str(source), source.stat().st_mtime)
    if _db_cache.get("key") != key:
        _db_cache["db"] = master_db_cls(str(source))
        _db_cache["key"] = key
    return _db_cache["db"], ("마스터 내장 고정값" if embedded else "기본 고정값")


def generate_document_from_master(master_path, kind, out_path):
    """공정마스터 정보로 kind 문서를 만들어 out_path(.xlsm)에 저장한다.

    (생성한 공정 수, 건너뛴 공정 수, 사용한 고정값 출처)를 반환한다."""
    make, builders, MasterDB = _load_engine()
    doc_name = KIND_TO_ENGINE_DOC[kind]

    data = read_master_project(master_path)
    if not data["rows"]:
        raise ValueError(f"'{MASTER_SHEET}' 시트에서 공정을 찾지 못했습니다. (공정번호 열이 비어 있습니다)")

    db, std_source = _get_db(MasterDB, master_path)

    prj = make.build_project(data, db)
    master_names = {(r["라인명"], r["공정번호"], r["공정Code"]): r["공정명"] for r in data["rows"]}
    for proc in prj.procs:
        name = master_names.get((proc["라인명"], proc["공정번호"], proc["공정Code"]))
        if name:
            proc["공정명"] = name
    if not prj.procs:
        raise ValueError("마스터의 공정Code가 고정값(STD)과 맞지 않아 만들 공정이 없습니다.")

    template = ENGINE_DIR / "templates" / make.TPL[doc_name]
    tmp_dir = tempfile.mkdtemp(prefix="pe_ams_gen_")
    try:
        produced = builders.generate(doc_name, str(template), db, prj, tmp_dir, None)
        Path(out_path).parent.mkdir(parents=True, exist_ok=True)
        shutil.move(produced, out_path)
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)
    return len(prj.procs), len(data["rows"]) - len(prj.procs), std_source
