"""공정 마스터 엑셀에서 '공법 설계'에 쓰는 고정값(JIG 원가 리스트, 설비 단가 리스트)을 뽑아 engine/method_seed.json 으로 저장한다.

사용법:  python tools/extract_method_seed.py "<공정마스터>.xlsm"
앱은 이 파일을 처음 한 번만 읽어서 고정값으로 복사한다. 이후에는 앱의 고정값 화면에서 고친다.
"""
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "engine" / "method_seed.json"

# JIG원가LIST 시트의 (이름 열, 가격 열) -> 고정값 목록 이름
JIG_COLUMNS = [("JIG BASE", 1), ("JIG TYPE", 3), ("차광막/안전센서", 5), ("EJECTOR", 7), ("JIG UNIT", 9), ("UTILITY", 11)]


def clean(value):
    return re.sub(r"\s*\n\s*", " ", str(value)).strip() if value is not None else ""


def price_of(value):
    return value if isinstance(value, (int, float)) else 0


def find_sheet(wb, *needles):
    for ws in wb:
        if all(n in ws.title for n in needles):
            return ws
    raise SystemExit(f"시트를 찾을 수 없습니다: {needles}")


def read_jig(wb):
    ws = find_sheet(wb, "JIG", "LIST") if any("JIG" in w.title and "LIST" in w.title for w in wb) else find_sheet(wb, "JIG")
    result = {}
    for label, col in JIG_COLUMNS:
        items = []
        for row in range(2, ws.max_row + 1):
            name = clean(ws.cell(row, col).value)
            if name:
                items.append({"name": name, "price": price_of(ws.cell(row, col + 1).value)})
        result[label] = items
    return result


def read_equipment(wb):
    ws = find_sheet(wb, "1.2")
    groups = {}
    kind = cat = None
    for row in range(7, ws.max_row + 1):
        b, c, d, e = (ws.cell(row, col).value for col in (2, 3, 4, 5))
        if b:
            kind = clean(b)
        if c:
            cat = clean(c)
        if d and kind and cat:
            groups.setdefault(f"{kind} · {cat}", []).append({"name": clean(d), "price": price_of(e)})
    return groups


def read_allowed(wb):
    """공정마다 고를 수 있는 설비 종류: 이름정의 '전용_용접_KEY_ASSY_용접' 같은 범위의 값들."""
    allowed = {}
    pattern = re.compile(r"^(전용|범용)_(용접|조립)_")
    for name, defined in wb.defined_names.items():
        if not pattern.match(name):
            continue
        for sheet_name, cells in defined.destinations:
            ws = wb[sheet_name]
            area = ws[cells.replace("$", "")]
            if not isinstance(area, tuple):  # 셀 1개
                area = ((area,),)
            values = [clean(c.value) for row in area for c in (row if isinstance(row, tuple) else (row,)) if c.value]
            allowed[name] = values
    return allowed


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    tmp = Path(tempfile.gettempdir()) / "method_seed_copy.xlsm"
    shutil.copy(sys.argv[1], tmp)
    wb = openpyxl.load_workbook(tmp, keep_vba=True)
    seed = {"jig": read_jig(wb), "equip": read_equipment(wb), "equip_allowed": read_allowed(wb)}
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(seed, ensure_ascii=False, indent=1), encoding="utf-8")
    print("JIG:", {k: len(v) for k, v in seed["jig"].items()})
    print("설비 종류:", len(seed["equip"]), " 설비 품목:", sum(len(v) for v in seed["equip"].values()))
    print("공정별 허용 설비 종류:", len(seed["equip_allowed"]))
    print("저장:", OUT)


if __name__ == "__main__":
    main()
